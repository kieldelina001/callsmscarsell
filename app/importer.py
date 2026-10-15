"""Read Column A (name) and Column B (mobile number) from .xlsx / .csv.

Only columns A and B are used. Rows keep their original order and their
original spreadsheet row number, so names and numbers can never be mixed up.
"""
import csv
import io
import os

from .phone_utils import SCI_RE, normalize_ph_mobile, number_to_text

HEADER_AUTO, HEADER_YES, HEADER_NO = "auto", "yes", "no"


# ---------- reading ----------
def read_raw_rows(path):
    """Return list of (excel_row_number, cellA, cellB)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return _read_xlsx(path)
    if ext == ".csv":
        return _read_csv(path)
    raise ValueError("Unsupported file type. Please choose a .xlsx or .csv file.")


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        rows = []
        for i, row in enumerate(ws.iter_rows(min_col=1, max_col=2, values_only=True), start=1):
            row = tuple(row) + (None,) * (2 - len(row))
            rows.append((i, row[0], row[1]))
        return rows
    finally:
        wb.close()


def _read_csv(path):
    with open(path, "rb") as f:
        data = f.read()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        a = row[0] if len(row) > 0 else ""
        b = row[1] if len(row) > 1 else ""
        rows.append((i, a, b))
    return rows


# ---------- cleaning / preview ----------
def _blank(v):
    return v is None or str(v).strip() == ""


def _clean_name(v):
    return "" if v is None else " ".join(str(v).split())


def _looks_like_header(number_cell):
    """Auto mode: header only if column B is NOT a number and has letters."""
    if normalize_ph_mobile(number_cell)["ok"]:
        return False
    s = number_to_text(number_cell)
    if not s or SCI_RE.match(s):
        return False
    return any(ch.isalpha() for ch in s)


def build_preview(raw_rows, header_mode=HEADER_AUTO):
    """Validate everything and return {'records': [...], 'summary': {...}}."""
    rows = [r for r in raw_rows if not (_blank(r[1]) and _blank(r[2]))]
    empty_skipped = len(raw_rows) - len(rows)

    header_row = None
    if rows:
        first = rows[0]
        if header_mode == HEADER_YES or (
                header_mode == HEADER_AUTO and _looks_like_header(first[2])):
            header_row = (_clean_name(first[1]), number_to_text(first[2]))
            rows = rows[1:]

    records, seen = [], {}
    for source_row, a, b in rows:
        name = _clean_name(a)
        norm = normalize_ph_mobile(b)
        issues = []
        if not name:
            issues.append("blank name")
        if not norm["ok"]:
            issues.append(norm["issue"])
        status = "invalid" if issues else "valid"
        if status == "valid":
            if norm["e164"] in seen:
                status = "duplicate"
                issues.append("duplicate of row %d" % seen[norm["e164"]])
            else:
                seen[norm["e164"]] = source_row
        records.append({
            "source_row": source_row,
            "name": name,
            "raw_number": number_to_text(b),
            "phone": norm["e164"],
            "local": norm["local"],
            "status": status,
            "issue": "; ".join(issues),
            "note": norm["note"],
        })

    count = lambda s: sum(1 for r in records if r["status"] == s)
    summary = {
        "data_rows": len(records),
        "valid": count("valid"),
        "invalid": count("invalid"),
        "duplicate": count("duplicate"),
        "empty_rows_skipped": empty_skipped,
        "header_skipped": header_row is not None,
        "header_row": header_row,
    }
    return {"records": records, "summary": summary}
